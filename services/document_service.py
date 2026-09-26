"""Saves transcripts/manuscripts to disk and renders Markdown -> DOCX.

The LLM only ever produces Markdown (spec 七: "不要直接让 LLM 输出
Word"). Turning that Markdown into a .docx is done here with a small,
dependency-light converter built on python-docx, so formatting stays
predictable and doesn't depend on the model behaving.
"""
from __future__ import annotations

import logging
import os
import re

from docx.oxml.ns import qn

logger = logging.getLogger("app.document")

# Unified fonts so Chinese and English characters render consistently.
# Heading colour is forced to black (the default python-docx Heading style
# is blue Calibri, which looks wrong next to Chinese body text).
BODY_FONT_CN = "微软雅黑"
BODY_FONT_EN = "Calibri"
HEADING_FONT_CN = "微软雅黑"
HEADING_FONT_EN = "Calibri"


def _set_run_font(run, cn_font: str, en_font: str, size_pt: float | None = None,
                  bold: bool | None = None, color_rgb: str | None = None):
    """Set both East-Asian and Latin fonts on a run so mixed CJK/English
    text renders in a consistent typeface instead of silently falling back."""
    run.font.name = en_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn_font)
    rFonts.set(qn("w:ascii"), en_font)
    rFonts.set(qn("w:hAnsi"), en_font)
    if size_pt is not None:
        run.font.size = None  # clear inherited
        from docx.shared import Pt
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if color_rgb is not None:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string(color_rgb)


def _style_doc_defaults(doc):
    """Override the default Normal / Heading styles so CJK and Latin share
    the same font family, and headings are black instead of python-docx's
    default blue."""
    # Normal (body text)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT_EN
    normal.font.size = None  # leave at default (~11pt)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT_CN)
    rfonts.set(qn("w:ascii"), BODY_FONT_EN)
    rfonts.set(qn("w:hAnsi"), BODY_FONT_EN)

    # Headings: black colour, same font family
    from docx.shared import RGBColor
    for level, size in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)):
        try:
            st = doc.styles[level]
        except KeyError:
            continue
        st.font.name = HEADING_FONT_EN
        st.font.size = None
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("000000")  # black
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            from docx.oxml import OxmlElement
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), HEADING_FONT_CN)
        rfonts.set(qn("w:ascii"), HEADING_FONT_EN)
        rfonts.set(qn("w:hAnsi"), HEADING_FONT_EN)


class DocumentError(RuntimeError):
    pass


def save_text(path: str, content: str, overwrite: bool = True) -> str:
    if os.path.exists(path) and not overwrite:
        logger.info("Not overwriting existing file: %s", path)
        return path
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp_path = path + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        f.write(content)
    os.replace(tmp_path, path)
    return path


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_markdown_paragraph(doc, line: str):
    """Very small Markdown -> docx renderer: headings, bold, plain text."""
    if line.startswith("### "):
        p = doc.add_heading(line[4:].strip(), level=3)
        for run in p.runs:
            _set_run_font(run, HEADING_FONT_CN, HEADING_FONT_EN, size_pt=13,
                          bold=True, color_rgb="000000")
        return
    if line.startswith("## "):
        p = doc.add_heading(line[3:].strip(), level=2)
        for run in p.runs:
            _set_run_font(run, HEADING_FONT_CN, HEADING_FONT_EN, size_pt=15,
                          bold=True, color_rgb="000000")
        return
    if line.startswith("# "):
        p = doc.add_heading(line[2:].strip(), level=1)
        for run in p.runs:
            _set_run_font(run, HEADING_FONT_CN, HEADING_FONT_EN, size_pt=18,
                          bold=True, color_rgb="000000")
        return

    p = doc.add_paragraph()
    pos = 0
    for m in _BOLD_RE.finditer(line):
        if m.start() > pos:
            run = p.add_run(line[pos:m.start()])
            _set_run_font(run, BODY_FONT_CN, BODY_FONT_EN)
        run = p.add_run(m.group(1))
        _set_run_font(run, BODY_FONT_CN, BODY_FONT_EN, bold=True)
        pos = m.end()
    if pos < len(line):
        run = p.add_run(line[pos:])
        _set_run_font(run, BODY_FONT_CN, BODY_FONT_EN)


def markdown_to_docx(markdown_text: str, docx_path: str) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise DocumentError(
            "生成 DOCX 需要 python-docx，请先运行: pip install python-docx"
        ) from e

    doc = docx.Document()
    _style_doc_defaults(doc)  # fix CJK / heading font + colour
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        _add_markdown_paragraph(doc, line)

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    try:
        doc.save(docx_path)
    except PermissionError as e:
        raise DocumentError(f"没有权限写入文件（文件可能被占用）: {docx_path}") from e
    except OSError as e:
        raise DocumentError(f"保存 DOCX 失败: {e}") from e
    return docx_path
